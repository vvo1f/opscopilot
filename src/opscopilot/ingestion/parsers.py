"""File -> list of (page number, text).

PDFs have real pages. DOCX and TXT don't, so they count as a single page.
Only text-based PDFs work: a scanned PDF is just images and would need OCR.
"""
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.table import Table
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".txt"}


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


@dataclass
class Page:
    number: int
    text: str


def _pdf_pages(path: Path) -> list[Page]:
    reader = PdfReader(str(path))
    return [Page(n, page.extract_text() or "") for n, page in enumerate(reader.pages, start=1)]


def _docx_pages(path: Path) -> list[Page]:
    doc = Document(str(path))
    lines = []
    # iter_inner_content yields paragraphs and tables in document order, so an
    # invoice table stays right after the text that introduces it
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        else:
            lines.append(block.text)
    return [Page(1, "\n".join(lines))]


def _txt_pages(path: Path) -> list[Page]:
    return [Page(1, path.read_text(encoding="utf-8", errors="replace"))]


PARSERS = {".pdf": _pdf_pages, ".docx": _docx_pages, ".txt": _txt_pages}


def parse_file(path: Path) -> list[Page]:
    suffix = path.suffix.lower()
    if suffix not in PARSERS:
        raise UnsupportedFileType(f"{suffix or 'no extension'} is not supported")
    pages = [p for p in PARSERS[suffix](path) if p.text.strip()]
    if not pages:
        raise EmptyDocument("no extractable text (a scanned PDF would need OCR)")
    return pages
